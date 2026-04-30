"""PPT 解析:每张幻灯片抽取标题+正文+备注,作为一"页"。"""
from __future__ import annotations
import hashlib
from pathlib import Path
from pptx import Presentation

from src.parsers.schema import ParsedDoc, PageBlock


def parse_pptx(path: Path) -> ParsedDoc:
    prs = Presentation(str(path))
    pages: list[PageBlock] = []

    for idx, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    parts.append(txt)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            parts.append(f"[备注] {notes}")

        text = "\n\n".join(parts)
        if not text:
            continue
        pages.append({
            "page": idx,
            "page_label": f"幻灯片 {idx}",
            "text": text,
            "tables": [],
        })

    return {
        "metainfo": {
            "file_name": path.name,
            "file_type": "pptx",
            "doc_id": _file_md5(path),
        },
        "content": pages,
    }


def _file_md5(path: Path) -> str:
    h = hashlib.md5()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()
